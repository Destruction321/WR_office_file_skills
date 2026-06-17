"""
# `.doc` / `.docx` 提取器。
- 先尝试 `python-docx`（可处理伪装成 `.doc` 的 `.docx`），再 COM 回退处理旧 `.doc` 格式。
"""

from pathlib import Path
from subprocess import run, DEVNULL, TimeoutExpired
from sys import platform

from .common import kill_orphan_com, DOC_SCRIPT
from .. import assets
from ..deps import ensure_import
from ..section import detect_chinese_heading
from ..util import safe_open_path, mktemp_in_dir

# 已知正文样式名（小写）—— 排除这些后，高频出现的自定义样式视为标题
_BODY_STYLE_NAMES = frozenset({
    'normal', '正文', '默认段落字体', 'body text', 'bodytext',
    'caption', 'footnote', 'endnote', 'header', 'footer',
    'toc 1', 'toc 2', 'toc 3', 'tocheading',
    'table grid', 'list paragraph',
})


def extract_doc(filepath: Path, assets_dir: Path | None = None) -> list[str]:
    """
    ## 提取旧格式 `.doc` 文件。
    
    - 先尝试 `python-docx`（可处理伪装成 `.doc` 的 `.docx`），再 COM 回退。
    - 如果 `python-docx` 只返回了错误信息（全是 `[Error...]` 或空），
    说明不是 `.docx` 变体，走 COM 路径。

    Args:
        filepath (Path): 文档文件路径。
        assets_dir (Path | None): 资源提取目标目录（可选）。

    Returns:
        lines (list[str]): 提取出的文本行，失败时返回错误信息。
    """
    try:
        result = extract_docx(filepath, assets_dir)
        if all(not line or line.startswith('[') for line in result):
            raise ValueError('python-docx 仅返回了错误信息')
        return result
    
    except Exception:
        return _extract_doc_com(filepath, assets_dir)


def extract_docx(filepath: Path, assets_dir: Path | None = None) -> list[str]:
    """
    - 通过 `python-docx` 提取 `.docx` 文件，保留段落/表格交错顺序和标题样式。

    Args:
        filepath (Path): 文档文件路径。
        assets_dir (Path | None): 资源提取目标目录（可选）。

    Returns:
        lines (list[str]): 提取出的文本行，失败时返回错误信息。
    """
    from docx.oxml.ns import qn

    lines: list[str] = []
    assets_result: dict[str, list[str]] = {}

    if assets_dir:
        assets_result = assets.extract_ooxml_assets(filepath, assets_dir / filepath.stem, '.docx')

    try:
        Document = ensure_import('python-docx', 'docx', attr='Document')  # type: ignore[assignment]

    except ImportError:
        return ['[Error: python-docx 未安装。执行: pip install python-docx]']

    with safe_open_path(filepath) as safe_path:
        try:
            doc = Document(str(safe_path))  # type: ignore[operator]

        except Exception as e:
            return [f'[Error: 用 python-docx 打开 .docx 失败: {e}]']

        # 预扫描：统计各样式出现次数，用于识别自定义标题样式
        style_count: dict[str, int] = {}
        for p_elem in doc.element.body.iter(qn("w:p")):
            pPr = p_elem.find(qn("w:pPr"))
            if pPr is None:
                continue
            pStyle = pPr.find(qn("w:pStyle"))
            if pStyle is None:
                continue
            val = pStyle.get(qn("w:val"), "")
            if val:
                style_count[val.lower()] = style_count.get(val.lower(), 0) + 1

        # 始终使用有序提取——保留段落/表格的真实交错顺序
        _extract_docx_body_ordered(doc, lines, style_count)

    if assets_result:
        assets.append_assets_summary(lines, assets_result)

    return lines


def _extract_doc_com(filepath: Path, assets_dir: Path | None = None) -> list[str]:
    """通过 Windows COM 提取旧 .doc 文件。"""
    if platform != 'win32':
        return ['[Error: 旧格式 .doc 提取需要 Windows + Microsoft Office]']

    if not Path(DOC_SCRIPT).exists():
        return ['[Error: 找不到 Word 提取脚本]']

    tmp_out = mktemp_in_dir(filepath, prefix='tmp_doc_') / 'output.txt'
    try:
        run(
            ['powershell', '-ExecutionPolicy', 'Bypass', '-File', DOC_SCRIPT,
             '-DocPath', str(filepath), '-OutFile', str(tmp_out)],
            stdout=DEVNULL, stderr=DEVNULL, timeout=120
        )
        if tmp_out.exists():
            lines = tmp_out.read_text(encoding='utf-8-sig').splitlines()
            if assets_dir:
                lines.append('[提示: 旧格式 .doc 暂不支持嵌入文件提取]')
            return lines
        return ['[Error: Word 提取未产生输出]']

    except TimeoutExpired:
        kill_orphan_com('WINWORD.EXE')
        return ['[Error: Word 提取超时]']

    except Exception as e:
        return [f'[Error: 通过 COM 提取 DOC 失败: {e}]']


def _extract_docx_body_ordered(doc, lines: list[str], style_count: dict[str, int]) -> None:
    """
    按文档顺序提取 docx 内容，添加 Markdown 标题标记。

    通过迭代 XML body 使段落和表格以真实文档顺序出现
    （有别于 doc.paragraphs + doc.tables 这种分离序列的方式）。
    """
    from docx.oxml.ns import qn

    def _text(elem) -> str:
        """收集 elem 内所有 w:t 节点文本。"""
        return ''.join(t.text or '' for t in elem.iter(qn('w:t')))

    body = doc.element.body
    for child in body:
        if child.tag == qn('w:p'):
            text = _text(child).strip()
            if not text:
                continue
            # 用多种方式检测标题级别，优先级：
            # 1. detect_chinese_heading 最准确（"实验七"=1、"实验目的"=2）
            # 2. 自定义样式推断（a4 → 2，与真实级别对比后可能高估或低估）
            # 取最小值（更高级别）作为最终级别
            heading_level = detect_chinese_heading(text)
            style_level = _get_heading_style_level(child, style_count)
            if style_level is not None and (heading_level is None or style_level < heading_level):
                heading_level = style_level
            if heading_level is not None:
                lines.append(f'{"#" * min(heading_level, 6)} {text}')
            else:
                lines.append(text)

        elif child.tag == qn('w:tbl'):
            lines.append('')
            for row_elem in child.iter(qn('w:tr')):
                cells = [
                    _text(tc).strip().replace('\n', ' ').replace('\r', '')
                    for tc in row_elem.iter(qn('w:tc'))
                ]
                if any(c for c in cells):
                    lines.append(' | '.join(cells))
            lines.append('')


def _get_heading_style_level(child, style_count: dict[str, int] | None = None) -> int | None:
    """
    检查段落是否应用了标题样式。

    按以下顺序检测：
    1. Word 内置标题样式（"Heading 1"-"Heading 9"）
    2. 自定义标题样式——通过文档中所有段落的样式出现频率推断：
       如果段落的样式不是正文样式，且在整个文档中 ≥2 次出现，则视为标题。

    没有匹配时返回 None。
    """
    from docx.oxml.ns import qn
    pPr = child.find(qn('w:pPr'))
    if pPr is None:
        return None

    pStyle = pPr.find(qn('w:pStyle'))
    if pStyle is None:
        return None

    style_val = pStyle.get(qn('w:val'), '')
    if not style_val:
        return None

    # 策略 1：Word 内置标题样式
    if style_val.lower().startswith('heading'):
        try:
            return int(style_val.split()[-1])
        except ValueError:
            return 1

    # 策略 2：通过样式名称在段落中出现的次数推断自定义标题样式
    # 高频出现的非正文样式很可能是自定义标题样式（如 a4, a3 等）
    if style_count and style_val.lower() not in _BODY_STYLE_NAMES:
        hits = style_count.get(style_val.lower(), 0)
        if hits >= 2:
            return 2  # 自定义标题默认视为 2 级

    return None
